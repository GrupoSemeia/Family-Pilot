"""Monta o documento de revisão de conteúdo (Word) para psicólogo/pediatra. Entrada: JSON de exportar.ts."""
import json, sys, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = json.load(open(sys.argv[1]))
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ('left_margin', 'right_margin'): setattr(sec, m, Cm(1.8))
sec.top_margin = sec.bottom_margin = Cm(1.6)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

def sombra(cell, cor):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), cor); tcPr.append(s)

def larguras(t, ws):
    t.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    for col, w in zip(t.columns, ws): col.width = Cm(w)
    for row in t.rows:
        for c, w in zip(row.cells, ws): c.width = Cm(w)

def repetir_cabecalho(row):
    p = row._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); p.append(h)

def par(txt, negrito=False, tam=None, cor=None, estilo=None, ital=False):
    p = doc.add_paragraph(style=estilo); r = p.add_run(txt); r.bold = negrito; r.italic = ital
    if tam: r.font.size = Pt(tam)
    if cor: r.font.color.rgb = RGBColor.from_string(cor)
    return p

class Mod:
    """Uma seção de tabelas de um módulo, com numeração estável das referências (ex.: SONO-014)."""
    def __init__(self, pfx): self.pfx, self.n = pfx, 0; self.t = None
    def tabela(self, titulo=None):
        if titulo: doc.add_heading(titulo, level=3)
        self.t = doc.add_table(rows=1, cols=3); self.t.style = 'Table Grid'
        for c, tx in zip(self.t.rows[0].cells, ('Ref.', 'Texto exibido ao usuário', 'Parecer (OK / ajustar / remover) e comentário')):
            c.text = ''; r = c.paragraphs[0].add_run(tx); r.bold = True; r.font.size = Pt(9.5); sombra(c, 'E8EEF4')
        repetir_cabecalho(self.t.rows[0])
    def linha(self, rotulo, texto):
        if not texto and texto != 0: return
        if self.t is None: self.tabela()
        self.n += 1
        cs = self.t.add_row().cells
        cs[0].text = ''; r = cs[0].paragraphs[0].add_run(f'{self.pfx}-{self.n:03d}'); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p = cs[1].paragraphs[0]
        if rotulo: b = p.add_run(rotulo + ': '); b.bold = True; b.font.size = Pt(10)
        p.add_run(str(texto)).font.size = Pt(10)
        cs[2].text = ''
        larguras(self.t, (2.0, 15.5, 8.6))
    def fim(self): self.t = None

def lista(M, rotulo, itens):
    for i in itens or []: M.linha(rotulo, i)

def frases(M, fr, faixas):
    for f in faixas:
        x = (fr or {}).get(f['id'])
        if x: M.linha(f"{f['rotulo']} · Acolher", x['a']); M.linha(f"{f['rotulo']} · Manter o limite", x['l'])

# ---------- capa e instruções ----------
t = doc.add_heading('FamilyPilot: revisão de conteúdo', 0)
par(f'Documento para o psicólogo(a) ou pediatra revisor(a). Gerado em {datetime.date.today().strftime("%d/%m/%Y")} a partir do conteúdo que está no aplicativo.', ital=True, cor='555555')
doc.add_heading('1. O que é o produto', 1)
par('O FamilyPilot é um aplicativo para pais, mães e responsáveis, usado principalmente pelo celular. Tem 7 módulos de orientação prática para o dia a dia com crianças e adolescentes: SOS Birra, Sono, Telas, Alimentação, Limites, Escola e Adolescência. Cada módulo traz leituras da situação, o que fazer e o que evitar, frases prontas para dizer ao filho (por faixa de idade), combinados, ideias, modelos de mensagem, registros e orientação sobre quando procurar ajuda profissional.')
par('Posicionamento: orientação educativa para pais. Não faz diagnóstico, não substitui pediatra, psicólogo ou escola. O app exibe esse aviso ao usuário.')
doc.add_heading('2. O que pedimos que você avalie', 1)
for x in [
    'Correção e segurança: alguma orientação pode ser prejudicial, desatualizada ou contrariar consensos da pediatria e da psicologia do desenvolvimento?',
    'Adequação à faixa etária: o que é dito para cada idade faz sentido (2 a 3, 4 a 5, 6 a 7, 8 a 10 anos; 11 a 12, 13 a 14, 15 a 17 no módulo Adolescência)?',
    'Linguagem: está clara, acolhedora e sem culpabilizar os pais? Há termos técnicos ou frases que soam como diagnóstico ou promessa de resultado?',
    'Sinais de alerta: os trechos "quando procurar ajuda" estão completos e no tom certo? Os serviços citados (telefones e recursos) estão corretos e atuais?',
    'Frases prontas: são respeitosas, realistas e seguras para dizer a uma criança naquela idade?',
    'Faltou algo importante? Há algum tema que não deveria estar aqui?']: doc.add_paragraph(x, style='List Bullet')
doc.add_heading('3. Como responder', 1)
par('Cada trecho tem uma referência (ex.: SONO-014). Na coluna da direita escreva OK, Ajustar ou Remover e, se quiser, o texto que sugere no lugar. Pode também usar os comentários do Word. Devolva o arquivo com as suas anotações; nós aplicamos os ajustes.')
par('As frases de cada situação aparecem por faixa de idade em duas versões: "Acolher" (reconhecer o sentimento) e "Manter o limite".', ital=True)

# ---------- módulos ----------
doc.add_page_break()
doc.add_heading('SOS Birra', 1)
b = d['birra']; cat = d['catalogo']
par(f"{cat['birra']['frase']}  {cat['birra']['detalhe']}", ital=True)
M = Mod('BIRRA')
faixas = b['IDADES']
M.tabela('Por idade: expectativa e dica')
for f in faixas:
    x = b['POR_IDADE'][f['id']]; M.linha(f"{f['rotulo']} · O que esperar", x['expectativa']); M.linha(f"{f['rotulo']} · Dica", x['dica'])
M.fim()
for mo in b['MOTIVOS']:
    M.tabela(f"Motivo da birra: {mo['nome']}")
    M.linha('Leitura', mo['leitura']); lista(M, 'Fazer', mo['fazer']); lista(M, 'Evitar', mo['evitar']); frases(M, mo['frases'], faixas); M.fim()
M.tabela('Segurança e o que evitar sempre')
lista(M, 'Segurança', b['SEGURANCA']); lista(M, 'Evitar sempre', b['EVITAR_SEMPRE']); lista(M, 'Depois da birra', b['DEPOIS']); M.fim()
for lo in b['LOCAIS']:
    M.tabela(f"Onde está: {lo['nome']}"); lista(M, 'Passo', lo['passos']); M.fim()
for c in b['COMPORTAMENTOS']:
    M.tabela(f"Comportamento: {c['nome']}"); frases(M, c['frases'], faixas); M.fim()
M.tabela('Gatilhos (Radar)')
for g in b['GATILHOS']: M.linha(g['nome'], g.get('dica'))
M.fim()
M.tabela('Opções do Diário (nomes exibidos)')
M.linha('Estratégias', '; '.join(e['nome'] for e in b['ESTRATEGIAS'])); M.linha('Resultados', '; '.join(e['nome'] for e in b['RESULTADOS'])); M.fim()

def ferramenta(M, fe, faixas):
    tp = fe['tipo']; M.tabela(f"{fe['icone']} {fe['nome']}"); par(fe['descricao'], ital=True, cor='555555')
    if tp == 'situacoes':
        M.fim()
        for s in fe['situacoes']:
            M.tabela(f"Situação: {s['nome']}"); M.linha('Leitura', s['leitura']); lista(M, 'Fazer', s['fazer']); lista(M, 'Evitar', s['evitar'])
            frases(M, s['frases'], faixas); lista(M, 'Depois', s['depois']); M.linha('Sinal de alerta', s.get('alerta')); M.fim()
        M.tabela('Por faixa de idade: o que esperar e dica')
        for f in faixas:
            x = fe['porFaixa'].get(f['id'])
            if x: M.linha(f"{f['rotulo']} · O que esperar", x['expectativa']); M.linha(f"{f['rotulo']} · Dica", x['dica'])
    elif tp == 'rotinaSono':
        for f in faixas:
            x = fe['porFaixa'].get(f['id'])
            if x: M.linha(f['rotulo'], f"dormir de {x['horas'][0]} a {x['horas'][1]} horas por noite; hora de dormir sugerida {x['dormir']}; acordar {x['acordar']}. {x['dica']}")
        for p in fe['passos']: M.linha('Passo da rotina' + (f" ({', '.join(p['faixas'])})" if p.get('faixas') else ''), f"{p['texto']} (começa {p['antes']} min antes de dormir)")
    elif tp == 'combinado':
        M.linha('Título', fe['titulo']); M.linha('Introdução', fe['intro']); M.linha('Fecho', fe['fecho'])
        for s in fe['secoes']:
            for it in s['itens']: M.linha(f"{s['titulo']}", it['texto'])
            for fx, tx in (s.get('dica') or {}).items(): M.linha(f"{s['titulo']} · dica {fx}", tx)
    elif tp == 'ideias':
        nomes = {c['id']: c['nome'] for c in fe['categorias']}
        for it in fe['itens']: M.linha(nomes.get(it['categoria'], it['categoria']) + (f" ({', '.join(it['faixas'])})" if it.get('faixas') else ''), it['texto'])
    elif tp == 'modelo':
        for mo in fe['modelos']: M.linha(f"Modelo: {mo['nome']}", mo['texto'])
    elif tp == 'registro':
        M.linha('Campos do registro', '; '.join(c['rotulo'] + (' (' + ', '.join(o['nome'] for o in c['opcoes']) + ')' if c.get('opcoes') else '') for c in fe['campos']))
    elif tp == 'ajuda':
        M.fim()
        for bl in fe['blocos']:
            M.tabela(bl['titulo']); lista(M, '', bl['itens']); M.fim()
        if fe.get('recursos'):
            M.tabela('Onde buscar ajuda'); [M.linha(r['nome'], r['texto']) for r in fe['recursos']]
    M.fim()

PFX = {'sono': 'SONO', 'telas': 'TELAS', 'alimentacao': 'ALIM', 'limites': 'LIM', 'escola': 'ESC', 'adolescencia': 'ADOL'}
total = M.n
for mid, mod in d['modulos'].items():
    doc.add_page_break(); doc.add_heading(f"{mod['icone']} {mod['nome']}", 1)
    par(f"{cat[mid]['frase']}  {cat[mid]['detalhe']}", ital=True); par(mod['resumo'])
    par('Faixas de idade: ' + ', '.join(f['rotulo'] for f in mod['faixas']), cor='555555')
    M = Mod(PFX[mid])
    for fe in mod['ferramentas']: ferramenta(M, fe, mod['faixas'])
    total += M.n

doc.add_page_break(); doc.add_heading('Parecer geral', 1)
par(f'Total de trechos neste documento: {total}.', cor='555555')
for q in ['Conteúdo aprovado para publicação?  ( ) Sim   ( ) Sim, com os ajustes marcados   ( ) Não',
          'Nome, registro profissional (CRP/CRM) e data da revisão:',
          'Temas que faltam ou que devem sair:',
          'Fontes ou diretrizes que recomenda citar no aplicativo (ex.: Sociedade Brasileira de Pediatria):',
          'Observações gerais:']:
    par(q, negrito=True); par('\n\n')
doc.save(sys.argv[2]); print('ok', total, 'trechos')
